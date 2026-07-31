"""QBO bank-rules loader for the brex-import skill's feed mode.

Accepts BOTH of:
  (a) the live `quickbooks_bank_rules_list` MCP output (forward-compat — the
      endpoint currently returns HTTP 400 "Unsupported Operation", but the
      JSON shape is documented in the spec and the parser is ready when QBO
      restores it), or
  (b) the .xls export from QBO Online → Banking → Rules → Export to Excel,
      converted to .xlsx first (the orchestrator does the conversion via
      `libreoffice --headless --convert-to xlsx`).

Both formats carry the same JSON inside — the .xls has columns Rule Name /
Rule Conditions / Rule Outputs where the latter two are JSON strings. The
live MCP returns parsed dicts. We normalize both into a common Rule shape.

Filtering (per spec):
  - Keep `ruleType 6` (description contains) — drop `ruleType 1` (description
    matches exact) — the 7 text-less rules in the production export are all
    type-1.
  - Keep direction `ruleType 10 == "-1"` (money out) — drop the 1 money-in
    (`Brex Dividend Received`).
  - Keep text length >= 3 — drop the 1 too-short rule.
  - `isAndRule` confirmed true across all 362 production rules.

Action types observed in the production export:
  0 = category (GL path string)
  1 = (unknown — present on 1 rule, ignored)
  2 = class (e.g. "110- G&A")
  5 = payee (vendor display name)
  7 = (unknown — present on a few rules, ignored)
  8 = auto-add (boolean)
  9 = (unknown)
 11 = (unknown — present on the money-in rule, ignored)

We extract 0/2/5. The rest are ignored without erroring.

The loader is pure-Python and does NO network. Caching to JSON happens in the
orchestrator layer (runner.py), not here.
"""

from __future__ import annotations
import json
import sys
from pathlib import Path
from dataclasses import dataclass, asdict


@dataclass
class Rule:
    name: str
    text: str          # the contains-this-text value (original case)
    text_lc: str       # lowercased for matching
    direction: str     # "-1" money-out
    payee: str | None
    category: str | None
    cls: str | None
    source: str        # "xls" or "live"

    def to_dict(self) -> dict:
        return asdict(self)


@dataclass
class SkipReason:
    name: str
    reason: str
    conds: dict | list


def _normalize_one(name: str, conds_json: dict, outs_json: dict,
                   source: str, skipped: list[SkipReason]) -> Rule | None:
    """Convert one rule's parsed JSON into a Rule, or return None + log if skipped."""
    conds = conds_json.get("ruleConditions", []) if isinstance(conds_json, dict) else []
    text_conds = [c for c in conds if c.get("ruleType") == 6]
    dir_conds = [c for c in conds if c.get("ruleType") == 10]

    if not text_conds:
        skipped.append(SkipReason(name, "no description-contains condition (ruleType 6)", conds))
        return None
    direction = None
    if dir_conds:
        direction = dir_conds[0].get("value")
    if direction != "-1":
        skipped.append(SkipReason(name, f"not money-out (direction={direction!r})", conds))
        return None

    text = (text_conds[0].get("value") or "").strip()
    if len(text) < 3:
        skipped.append(SkipReason(name, f"text too short ({text!r}) — false-positive risk", conds))
        return None

    # Pull actions
    actions = outs_json.get("ruleActions", []) if isinstance(outs_json, dict) else []
    by_type = {a.get("actionType"): a.get("value") for a in actions}
    return Rule(
        name=name,
        text=text,
        text_lc=text.lower(),
        direction=direction,
        payee=by_type.get(5),
        category=by_type.get(0),
        cls=by_type.get(2),
        source=source,
    )


def load_from_xlsx(path: str | Path) -> tuple[list[Rule], list[SkipReason]]:
    """Load rules from a Banking → Rules → Export-to-Excel workbook.

    The QBO export is .xls (Excel 97-2003). openpyxl can't read .xls, so the
    orchestrator pre-converts to .xlsx via `libreoffice --headless --convert-to
    xlsx`. This function expects the already-converted .xlsx.
    """
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.active

    headers = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(1, c).value
        if v:
            headers[str(v).strip().lower()] = c
    name_col = headers.get("rule name")
    cond_col = headers.get("rule conditions")
    out_col = headers.get("rule outputs")
    if not (name_col and cond_col and out_col):
        raise ValueError(
            f"qbo_rules.xlsx is missing expected headers. Got: {list(headers.keys())!r}. "
            "Expected 'Rule Name', 'Rule Conditions', 'Rule Outputs'."
        )

    rules: list[Rule] = []
    skipped: list[SkipReason] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        name = row[name_col - 1]
        cond_raw = row[cond_col - 1]
        out_raw = row[out_col - 1]
        if not name:
            continue
        try:
            conds = json.loads(cond_raw) if cond_raw else {}
            outs = json.loads(out_raw) if out_raw else {}
        except (json.JSONDecodeError, TypeError) as e:
            skipped.append(SkipReason(str(name), f"JSON parse fail: {e}", {}))
            continue
        rule = _normalize_one(str(name), conds, outs, "xls", skipped)
        if rule:
            rules.append(rule)
    wb.close()
    rules.sort(key=lambda r: -len(r.text_lc))
    return rules, skipped


def load_from_live(payload: list[dict]) -> tuple[list[Rule], list[SkipReason]]:
    """Load rules from the live `quickbooks_bank_rules_list` MCP output.

    The live endpoint currently returns HTTP 400; this is for forward-compat.
    Expected shape (best-effort per QBO docs):
        [{"Name": "...", "Conditions": {...}, "Actions": {...}}, ...]
    When QBO restores the endpoint we'll need to confirm exact key names and
    adjust here. For now we accept the most likely variants.
    """
    rules: list[Rule] = []
    skipped: list[SkipReason] = []
    for r in payload:
        name = r.get("Name") or r.get("name") or r.get("RuleName") or ""
        conds = (r.get("Conditions") or r.get("conditions")
                 or {"ruleConditions": r.get("ruleConditions", [])}
                 or r.get("Rule Conditions"))
        outs = (r.get("Actions") or r.get("actions")
                or {"ruleActions": r.get("ruleActions", [])}
                or r.get("Rule Outputs"))
        if isinstance(conds, str):
            conds = json.loads(conds)
        if isinstance(outs, str):
            outs = json.loads(outs)
        rule = _normalize_one(str(name), conds or {}, outs or {}, "live", skipped)
        if rule:
            rules.append(rule)
    rules.sort(key=lambda r: -len(r.text_lc))
    return rules, skipped


def load_rules(path_or_payload) -> tuple[list[Rule], list[SkipReason]]:
    """Dispatch on input type: a Path/str → .xlsx file; a list → live payload."""
    if isinstance(path_or_payload, (str, Path)):
        return load_from_xlsx(path_or_payload)
    if isinstance(path_or_payload, list):
        return load_from_live(path_or_payload)
    raise TypeError(f"load_rules: expected path or list, got {type(path_or_payload).__name__}")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: rules_loader.py <path-to-qbo_rules.xlsx>", file=sys.stderr)
        sys.exit(2)
    rules, skipped = load_from_xlsx(sys.argv[1])
    skip_buckets: dict[str, int] = {}
    for r in skipped:
        key = r.reason.split(" ")[0:4]
        k = " ".join(key)[:60]
        skip_buckets[k] = skip_buckets.get(k, 0) + 1
    summary = dict(
        kept=len(rules),
        skipped=len(skipped),
        skip_reasons=skip_buckets,
        longest_text=max((len(r.text_lc) for r in rules), default=0),
        shortest_text=min((len(r.text_lc) for r in rules), default=0),
        sample=[(r.name, r.text, r.payee, r.category, r.cls) for r in rules[:5]],
    )
    print(json.dumps(summary, indent=2, default=str))
