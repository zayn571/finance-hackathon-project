#!/usr/bin/env python3
"""
Benchmark for build_department_listing.py — baseline vs optimized.

Generates synthetic BambooHR report JSON at several sizes, profiles each
implementation, and reports per-section timings and speedup.

Usage:
  python benchmark.py
"""
import sys
import io
import pstats
import cProfile
import tempfile
import shutil
import datetime
import time
from pathlib import Path
from copy import copy

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import openpyxl
from openpyxl.worksheet.formula import ArrayFormula

# Import the new helpers from the optimized build script
from build_department_listing import (
    _get_style_ids_from_xml, _find_dl_zip_entry, _extract_header_row_xml,
    _build_sheetdata_xml, _replace_sheetdata,
    parse_date, prefix_travel, prefix_payroll,
    REFERENCE,
)
import zipfile

SKILL_DIR = HERE.parent

# ── synthetic data ────────────────────────────────────────────────────────────

TEAMS = [
    "100- G&A:110- Finance",
    "100- G&A:120- HR",
    "200- Sales:210- AEs",
    "300- ServiceNow:320- Engineering",
    "400- Cloud:410- DevOps",
]
STATES    = ["MA", "NY", "CA", "TX", "WA", "FL"]
LOCS      = ["Boston, MA", "New York, NY", "San Francisco, CA"]
DEPTS     = ["Finance", "Engineering", "Sales", "HR", "DevOps"]
DIVS      = ["G&A", "Engineering", "Revenue"]
LEVELS    = ["Junior", "Mid", "Senior", "Staff", "Principal"]


def make_employee(i):
    return {
        "employeeNumber": str(i),
        "firstName": "First%d" % i,
        "lastName": "Last%d" % i,
        "fullName2": "Last%d, First%d" % (i, i),
        "gender": "Male" if i % 2 == 0 else "Female",
        "hireDate": "2020-%02d-%02d" % ((i % 12) + 1, (i % 28) + 1),
        "employmentHistoryStatus": "Active",
        "employeeStatusDate": "2020-%02d-%02d" % ((i % 12) + 1, (i % 28) + 1),
        "division": DIVS[i % len(DIVS)],
        "department": DEPTS[i % len(DEPTS)],
        "location": LOCS[i % len(LOCS)],
        "jobTitle": "Engineer %d" % i,
        "teams": TEAMS[i % len(TEAMS)],
        "customBillable": "Yes" if i % 3 == 0 else "No",
        "customEmployeeLevel": LEVELS[i % len(LEVELS)],
        "state": STATES[i % len(STATES)],
    }


def make_roster(n):
    return [make_employee(i) for i in range(1, n + 1)]


# ── baseline implementation (original) ───────────────────────────────────────

def build_baseline(roster, outpath):
    shutil.copyfile(REFERENCE, outpath)
    wb = openpyxl.load_workbook(outpath, data_only=False)
    ws = wb["Department Listing"]

    tmpl = {c: copy(ws.cell(2, c)._style) for c in range(1, 20)}

    for r in range(2, ws.max_row + 1):
        for c in range(1, 20):
            ws.cell(r, c).value = None

    r = 2
    for e in roster:
        num = int(e["employeeNumber"])
        teams = e.get("teams")
        ws.cell(r, 1).value = num
        ws.cell(r, 2).value = e.get("firstName")
        ws.cell(r, 3).value = e.get("lastName")
        ws.cell(r, 4).value = e.get("fullName2")
        ws.cell(r, 5).value = e.get("gender")
        ws.cell(r, 6).value = parse_date(e.get("hireDate"))
        ws.cell(r, 7).value = e.get("employmentHistoryStatus")
        ws.cell(r, 8).value = parse_date(e.get("employeeStatusDate"))
        ws.cell(r, 9).value = e.get("division")
        ws.cell(r, 10).value = e.get("department")
        ws.cell(r, 11).value = e.get("location")
        ws.cell(r, 12).value = e.get("jobTitle")
        ws.cell(r, 13).value = teams
        ws.cell(r, 14).value = e.get("customBillable")
        ws.cell(r, 15).value = e.get("customEmployeeLevel")
        ws.cell(r, 16).value = e.get("state")
        ws.cell(r, 17).value = "=M%d" % r
        ws.cell(r, 18).value = ArrayFormula("R%d" % r, "=" + prefix_travel(r))
        ws.cell(r, 19).value = ArrayFormula("S%d" % r, "=" + prefix_payroll(r))
        for c in range(1, 20):
            ws.cell(r, c)._style = tmpl[c]
        r += 1

    wb.save(outpath)


# ── optimized implementation ──────────────────────────────────────────────────

def build_optimized(roster, outpath):
    dl_zip_entry   = _find_dl_zip_entry(REFERENCE)
    with zipfile.ZipFile(REFERENCE) as zf:
        dl_xml_bytes = zf.read(dl_zip_entry)
    style_ids      = _get_style_ids_from_xml(dl_xml_bytes)
    header_row_xml = _extract_header_row_xml(dl_xml_bytes)
    new_sheetdata  = _build_sheetdata_xml(roster, style_ids, header_row_xml)
    new_dl_bytes   = _replace_sheetdata(dl_xml_bytes, new_sheetdata)

    with zipfile.ZipFile(REFERENCE, "r") as src, \
         zipfile.ZipFile(outpath, "w", compression=zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            dst.writestr(item, new_dl_bytes if item.filename == dl_zip_entry
                         else src.read(item.filename))


# ── timing helpers ────────────────────────────────────────────────────────────

SIZES = [10, 50, 100, 250, 500]
REPEATS = 3


def time_build(fn, roster, tmpdir, tag, idx):
    times = []
    for _ in range(REPEATS):
        out = tmpdir / ("%s_%d_%d.xlsx" % (tag, len(roster), idx))
        if out.exists():
            out.unlink()
        t0 = time.perf_counter()
        fn(roster, out)
        times.append(time.perf_counter() - t0)
    return min(times)   # best of N to reduce OS noise


def main():
    print("=== BambooHR Department Listing — Benchmark (baseline vs optimised) ===\n")

    if not REFERENCE.exists():
        print("ERROR: reference workbook not found at", REFERENCE)
        sys.exit(1)

    tmpdir = Path(tempfile.mkdtemp(prefix="bldl_bench_"))
    try:
        # ── 1. Wall-clock comparison across sizes ─────────────────────────────
        print("--- Wall-clock time (best of %d runs, seconds) ---" % REPEATS)
        print("%-6s  %-10s  %-10s  %-8s" % ("N", "baseline", "optimised", "speedup"))
        print("-" * 44)

        for n in SIZES:
            roster = make_roster(n)
            t_base = time_build(build_baseline,  roster, tmpdir, "base", n)
            t_opt  = time_build(build_optimized, roster, tmpdir, "opt",  n)
            speedup = t_base / t_opt if t_opt > 0 else float("inf")
            print("%-6d  %-10.3f  %-10.3f  %-8.1fx" % (n, t_base, t_opt, speedup))

        print()

        # ── 2. cProfile on both at largest size ───────────────────────────────
        largest = SIZES[-1]
        roster_large = make_roster(largest)

        for label, fn in [("baseline", build_baseline), ("optimised", build_optimized)]:
            print("--- cProfile top-15 hotspots: %s (N=%d) ---" % (label, largest))
            out = tmpdir / ("profile_%s.xlsx" % label)
            if out.exists():
                out.unlink()
            pr = cProfile.Profile()
            pr.enable()
            fn(roster_large, out)
            pr.disable()
            s = io.StringIO()
            pstats.Stats(pr, stream=s).sort_stats("cumulative").print_stats(15)
            print(s.getvalue())

        # ── 3. Scaling analysis ───────────────────────────────────────────────
        print("--- Scaling: optimised time per row (ms) ---")
        print("%-6s  %-12s" % ("N", "ms/row"))
        for n in SIZES:
            roster = make_roster(n)
            t = time_build(build_optimized, roster, tmpdir, "scale", n)
            print("%-6d  %-12.3f" % (n, 1000 * t / n))

    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)


if __name__ == "__main__":
    main()
